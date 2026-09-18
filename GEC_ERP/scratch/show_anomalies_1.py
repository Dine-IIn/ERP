import sys, os
sys.stdout.reconfigure(encoding='utf-8')
import pandas as pd
import openpyxl
from collections import defaultdict

wb = openpyxl.load_workbook('1-ERP 90-200 TON.xlsx', data_only=True)

all_items = []
for sname in wb.sheetnames:
    ws = wb[sname]
    max_r = ws.max_row
    max_c = ws.max_column
    
    row1 = [cell.value for cell in list(ws.iter_rows(min_row=1, max_row=1))[0]]
    row2 = [cell.value for cell in list(ws.iter_rows(min_row=2, max_row=2))[0]]
    
    col_idx = 0
    while col_idx < max_c:
        h1 = row1[col_idx]
        h2 = row2[col_idx] if col_idx < len(row2) else None
        
        if h2 and 'SR' in str(h2).upper():
            block_title = h1
            if not block_title:
                for b_idx in range(col_idx, -1, -1):
                    if row1[b_idx]:
                        block_title = row1[b_idx]
                        break
            
            subheaders = []
            for c in range(col_idx, max_c):
                val = row2[c]
                if c > col_idx and (val is None or (row1[c] and 'SR' not in str(val).upper() and row1[c] != block_title)):
                    break
                subheaders.append(str(val).strip().replace('\n', ' ') if val is not None else '')
                if 'CLASS' in subheaders[-1].upper():
                    break
            
            for r in range(3, max_r + 1):
                row_cells = [ws.cell(row=r, column=col_idx + 1 + offset).value for offset in range(len(subheaders))]
                sr_no = row_cells[0]
                if sr_no is None or str(sr_no).strip() == '':
                    continue
                try:
                    sr_num = int(float(str(sr_no).strip()))
                except:
                    continue
                
                row_dict = {'sheet': sname, 'block': str(block_title).strip(), 'row': r, 'sr_no': sr_num}
                for sh, cell_val in zip(subheaders, row_cells):
                    sh_clean = sh.upper()
                    if 'OLD' in sh_clean:
                        row_dict['old_code'] = str(cell_val).strip() if cell_val is not None else ''
                    elif 'NEW' in sh_clean or 'PART CODE' in sh_clean:
                        row_dict['new_code'] = str(cell_val).strip() if cell_val is not None else ''
                    elif 'DECRIPTION' in sh_clean or 'DESCRIPTION' in sh_clean or 'ITEM' in sh_clean:
                        row_dict['name'] = str(cell_val).strip() if cell_val is not None else ''
                    elif 'QTY' in sh_clean:
                        row_dict['qty'] = float(cell_val) if cell_val is not None and str(cell_val).strip() != '' else 1.0
                    elif 'CLASS' in sh_clean:
                        row_dict['class'] = str(cell_val).strip() if cell_val is not None else ''
                
                all_items.append(row_dict)
            col_idx += max(1, len(subheaders))
        else:
            col_idx += 1

wb.close()

# Anomaly 1: Same Part Code with Different Descriptions
code_to_names = defaultdict(lambda: defaultdict(list))
for it in all_items:
    c = it.get('new_code', '')
    n = it.get('name', '')
    if c and c != '-' and n and n != '-':
        code_to_names[c][n].append(f"{it['sheet']} > {it['block']} (Row {it['row']})")

print("=== ALL ANOMALY 1: SAME NEW PART CODE WITH DIFFERENT DESCRIPTIONS ===")
for c, names_dict in sorted(code_to_names.items()):
    if len(names_dict) > 1:
        print(f"\nPart Code: '{c}'")
        for name, occurrences in names_dict.items():
            print(f"  - Description: \"{name}\"")
            for occ in occurrences:
                print(f"      {occ}")
