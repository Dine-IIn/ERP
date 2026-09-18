import sys, os
sys.stdout.reconfigure(encoding='utf-8')
import pandas as pd
import openpyxl

wb = openpyxl.load_workbook('1-ERP 90-200 TON.xlsx', data_only=True)

print("=== SHEETS IN 1-ERP 90-200 TON.xlsx ===")
for sname in wb.sheetnames:
    ws = wb[sname]
    print(f"\nSheet: '{sname}' (Rows: {ws.max_row}, Cols: {ws.max_column})")
    row1 = [cell.value for cell in list(ws.iter_rows(min_row=1, max_row=1))[0]]
    row2 = [cell.value for cell in list(ws.iter_rows(min_row=2, max_row=2))[0]]
    for idx, (h1, h2) in enumerate(zip(row1, row2)):
        if h1 or (h2 and 'SR' in str(h2).upper()):
            print(f"  Col {idx+1}: Header='{h1}' | SubHeader='{h2}'")

wb.close()
