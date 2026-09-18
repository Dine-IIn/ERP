import sys, os
sys.stdout.reconfigure(encoding='utf-8')
import pandas as pd
import openpyxl

wb = openpyxl.load_workbook('1-ERP 90-200 TON.xlsx', data_only=True)
ws = wb['CASTING-90-200 TON']
print("=== CASTING-90-200 TON ===")
for r in range(1, ws.max_row + 1):
    vals = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
    # filter out None if all None
    if any(vals):
        print(f"Row {r:2d}:", [v for v in vals if v is not None][:10])

print("\n=== PATTERN-90-200 TON ===")
ws_p = wb['PATTERN-90-200 TON']
for r in range(1, ws_p.max_row + 1):
    vals = [ws_p.cell(row=r, column=c).value for c in range(1, ws_p.max_column + 1)]
    if any(vals):
        print(f"Row {r:2d}:", [v for v in vals if v is not None][:10])

wb.close()
