import sys, os
sys.stdout.reconfigure(encoding='utf-8')
import openpyxl
from collections import defaultdict

wb = openpyxl.load_workbook('1-ERP BOM 90-200 TON.xlsx', data_only=True)
all_items = []

for sname in wb.sheetnames:
    ws = wb[sname]
    max_r = ws.max_row
    max_c = ws.max_column
    
    row1 = [cell.value for cell in list(ws.iter_rows(min_row=1, max_row=1))[0]]
    row2 = [cell.value for cell in list(ws.iter_rows(min_row=2, max_row=2))[0]]
    
    if sname == 'Valve 90-200 TON':
        for col_start, blk in [(1, 'VALVE 90/110 TON'), (6, 'VALVE 140/200 TON')]:
            for r in range(5, max_r + 1):
                sr = ws.cell(r, col_start+1).value
                desc = ws.cell(r, col_start+2).value
                qty = ws.cell(r, col_start+3).value
                sub_unit = ws.cell(r, col_start+4).value
                if desc and str(desc).strip():
                    all_items.append({
                        'sheet': sname,
                        'block': blk,
                        'row': r,
                        'old_code': '',
                        'new_code': '',
                        'name': str(desc).strip(),
                        'qty': float(qty) if qty and str(qty).replace('.','').isdigit() else 1.0,
                        'class': 'HY'
                    })
        continue

    col_idx = 0
    while col_idx < max_c:
        h1 = row1[col_idx]
        h2 = row2[col_idx] if col_idx < len(row2) else None
        
        if h2 and 'SR' in str(h2).upper():
            block_title = h1
            if not block_title:
                for b_idx in range(col_idx, -1, -1):
                    if row1[b_idx]:
                        block_title = row1[b_idx]
                        break
            
            subheaders = []
            for c in range(col_idx, max_c):
                val = row2[c]
                if c > col_idx and (val is None or (row1[c] and 'SR' not in str(val).upper() and row1[c] != block_title)):
                    break
                subheaders.append(str(val).strip().replace('\n', ' ') if val is not None else '')
                if 'CLASS' in subheaders[-1].upper():
                    break
            
            for r in range(3, max_r + 1):
                row_cells = [ws.cell(row=r, column=col_idx + 1 + offset).value for offset in range(len(subheaders))]
                sr_no = row_cells[0]
                if sr_no is None or str(sr_no).strip() == '':
                    continue
                try:
                    sr_num = int(float(str(sr_no).strip()))
                except:
                    continue
                
                row_dict = {'sheet': sname, 'block': str(block_title).strip(), 'row': r, 'sr_no': sr_num}
                for sh, cell_val in zip(subheaders, row_cells):
                    sh_clean = sh.upper()
                    if 'OLD' in sh_clean:
                        row_dict['old_code'] = str(cell_val).strip() if cell_val is not None else ''
                    elif 'NEW' in sh_clean or 'PART CODE' in sh_clean:
                        row_dict['new_code'] = str(cell_val).strip() if cell_val is not None else ''
                    elif 'DECRIPTION' in sh_clean or 'DESCRIPTION' in sh_clean or 'ITEM' in sh_clean:
                        row_dict['name'] = str(cell_val).strip() if cell_val is not None else ''
                    elif 'QTY' in sh_clean:
                        row_dict['qty'] = float(cell_val) if cell_val is not None and str(cell_val).strip() != '' else 1.0
                    elif 'CLASS' in sh_clean:
                        row_dict['class'] = str(cell_val).strip() if cell_val is not None else ''
                
                all_items.append(row_dict)
            col_idx += max(1, len(subheaders))
        else:
            col_idx += 1

wb.close()

print(f"Total rows parsed: {len(all_items)}")

# Anomaly 1: Same Part Code with Different Descriptions
code_to_names = defaultdict(lambda: defaultdict(list))
for it in all_items:
    c = str(it.get('new_code', '')).strip()
    n = str(it.get('name', '')).strip()
    if c and c != '-' and n and n != '-':
        code_to_names[c][n].append(f"{it['sheet']} > {it['block']} (Row {it['row']})")

print("\n" + "="*80)
print("ANOMALY REPORT 1: SAME PART CODE WITH DIFFERENT DESCRIPTIONS")
print("="*80)
count_a1 = 0
for c, names_dict in sorted(code_to_names.items()):
    if len(names_dict) > 1:
        count_a1 += 1
        print(f"\n[!] Part Code: {c}")
        for name, occurrences in names_dict.items():
            print(f"    - \"{name}\" ({len(occurrences)} times) e.g. {occurrences[0]}")

print(f"\n>>> Total conflicting Part Codes: {count_a1}")

# Anomaly 2: Same Description with Different Part Codes
name_to_codes = defaultdict(lambda: defaultdict(list))
for it in all_items:
    c = str(it.get('new_code', '')).strip()
    n = str(it.get('name', '')).strip()
    if c and c != '-' and n and n != '-':
        name_to_codes[n][c].append(f"{it['sheet']} > {it['block']} (Row {it['row']})")

print("\n" + "="*80)
print("ANOMALY REPORT 2: SAME DESCRIPTION WITH DIFFERENT PART CODES")
print("="*80)
count_a2 = 0
for n, codes_dict in sorted(name_to_codes.items()):
    if len(codes_dict) > 1:
        count_a2 += 1
        print(f"\n[!] Description: \"{n}\"")
        for code, occurrences in codes_dict.items():
            print(f"    - Part Code: {code} ({len(occurrences)} times) e.g. {occurrences[0]}")

print(f"\n>>> Total conflicting Descriptions: {count_a2}")
