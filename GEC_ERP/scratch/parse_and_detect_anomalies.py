import sys, os
sys.stdout.reconfigure(encoding='utf-8')
import pandas as pd
import openpyxl
from collections import defaultdict

wb = openpyxl.load_workbook('1-ERP 90-200 TON.xlsx', data_only=True)

all_items = []
bom_records = defaultdict(list)
raw_castings = {} # finished_code -> raw casting info or pattern info
patterns = {}

for sname in wb.sheetnames:
    ws = wb[sname]
    max_r = ws.max_row
    max_c = ws.max_column
    
    # scan for block headers in row 1
    row1 = [cell.value for cell in list(ws.iter_rows(min_row=1, max_row=1))[0]]
    row2 = [cell.value for cell in list(ws.iter_rows(min_row=2, max_row=2))[0]]
    
    col_idx = 0
    while col_idx < max_c:
        h1 = row1[col_idx]
        h2 = row2[col_idx] if col_idx < len(row2) else None
        
        # Check if this column starts a table block
        # Usually SR NO is in col_idx
        if h2 and 'SR' in str(h2).upper():
            # Determine block title from h1 or backward non-empty cell
            block_title = h1
            if not block_title:
                # search backward in row1
                for b_idx in range(col_idx, -1, -1):
                    if row1[b_idx]:
                        block_title = row1[b_idx]
                        break
            
            # Subheaders in row 2 starting at col_idx
            subheaders = []
            block_cols = 0
            for c in range(col_idx, max_c):
                val = row2[c]
                if c > col_idx and (val is None or (row1[c] and 'SR' not in str(val).upper() and row1[c] != block_title)):
                    # Next block started or blank separator
                    break
                subheaders.append(str(val).strip().replace('\n', ' ') if val is not None else '')
                block_cols += 1
                if 'CLASS' in subheaders[-1].upper():
                    break
            
            # Read rows for this block
            # print(f"Sheet '{sname}' Block '{block_title}' from col {col_idx+1} ({len(subheaders)} cols): {subheaders}")
            for r in range(3, max_r + 1):
                row_cells = [ws.cell(row=r, column=col_idx + 1 + offset).value for offset in range(len(subheaders))]
                sr_no = row_cells[0]
                if sr_no is None or str(sr_no).strip() == '':
                    continue
                try:
                    # check if sr_no is number-like
                    sr_num = int(float(str(sr_no).strip()))
                except:
                    continue
                
                # Extract columns based on subheaders
                row_dict = {'sheet': sname, 'block': str(block_title).strip(), 'row': r}
                for sh, cell_val in zip(subheaders, row_cells):
                    sh_clean = sh.upper()
                    if 'OLD' in sh_clean:
                        row_dict['old_code'] = str(cell_val).strip() if cell_val is not None else ''
                    elif 'NEW' in sh_clean or 'PART CODE' in sh_clean:
                        row_dict['new_code'] = str(cell_val).strip() if cell_val is not None else ''
                    elif 'DECRIPTION' in sh_clean or 'DESCRIPTION' in sh_clean or 'ITEM' in sh_clean:
                        row_dict['name'] = str(cell_val).strip() if cell_val is not None else ''
                    elif 'QTY' in sh_clean:
                        row_dict['qty'] = float(cell_val) if cell_val is not None and str(cell_val).strip() != '' else 1.0
                    elif 'CLASS' in sh_clean:
                        row_dict['class'] = str(cell_val).strip() if cell_val is not None else ''
                
                all_items.append(row_dict)
            
            col_idx += max(1, len(subheaders))
        else:
            col_idx += 1

wb.close()

df = pd.DataFrame(all_items)
print(f"Total extracted rows across all sheets: {len(df)}")
print(df.head(10))

# ANOMALY DETECTION
print("\n" + "="*50)
print("=== ANOMALY CHECK 1: Same NEW_CODE with different ITEM NAMES ===")
code_to_names = defaultdict(set)
for it in all_items:
    c = it.get('new_code', '')
    n = it.get('name', '')
    if c and c != '-' and n and n != '-':
        code_to_names[c].add(n)

anomaly_1_count = 0
for c, names in code_to_names.items():
    if len(names) > 1:
        anomaly_1_count += 1
        print(f"  [ANOMALY 1] Code '{c}' has {len(names)} different descriptions:")
        for n in names:
            # show where it appeared
            occ = [f"{x['sheet']} / {x['block']}" for x in all_items if x.get('new_code') == c and x.get('name') == n]
            print(f"    - \"{n}\" in: {occ[:2]}")

if anomaly_1_count == 0:
    print("  ✓ No anomalies found for Same Code -> Different Names.")

print("\n" + "="*50)
print("=== ANOMALY CHECK 2: Same ITEM NAME with different NEW_CODES ===")
name_to_codes = defaultdict(set)
for it in all_items:
    c = it.get('new_code', '')
    n = it.get('name', '')
    if c and c != '-' and n and n != '-':
        name_to_codes[n.upper()].add(c)

anomaly_2_count = 0
for n, codes in name_to_codes.items():
    if len(codes) > 1:
        anomaly_2_count += 1
        print(f"  [ANOMALY 2] Item \"{n}\" has {len(codes)} different new part codes:")
        for c in codes:
            occ = [f"{x['sheet']} / {x['block']}" for x in all_items if x.get('new_code') == c and x.get('name', '').upper() == n]
            print(f"    - Code '{c}' in: {occ[:2]}")

if anomaly_2_count == 0:
    print("  ✓ No anomalies found for Same Name -> Different Codes.")

print("\n" + "="*50)
print("=== ANOMALY CHECK 3: Same OLD_CODE with different NEW_CODES ===")
old_to_new = defaultdict(set)
for it in all_items:
    o = it.get('old_code', '')
    c = it.get('new_code', '')
    if o and o != '-' and c and c != '-':
        old_to_new[o].add(c)

anomaly_3_count = 0
for o, codes in old_to_new.items():
    if len(codes) > 1:
        anomaly_3_count += 1
        print(f"  [ANOMALY 3] Old Code '{o}' mapped to multiple new codes: {codes}")

if anomaly_3_count == 0:
    print("  ✓ No anomalies found for Same Old Code -> Different New Codes.")

print("\n" + "="*50)
print("=== ANOMALY CHECK 4: Same NEW_CODE with different OLD_CODES ===")
new_to_old = defaultdict(set)
for it in all_items:
    o = it.get('old_code', '')
    c = it.get('new_code', '')
    if o and o != '-' and c and c != '-':
        new_to_old[c].add(o)

anomaly_4_count = 0
for c, olds in new_to_old.items():
    if len(olds) > 1:
        anomaly_4_count += 1
        print(f"  [ANOMALY 4] New Code '{c}' mapped to multiple old codes: {olds}")

if anomaly_4_count == 0:
    print("  ✓ No anomalies found for Same New Code -> Different Old Codes.")
