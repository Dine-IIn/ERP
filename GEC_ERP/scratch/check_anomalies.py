import sys, os
sys.stdout.reconfigure(encoding='utf-8')
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from collections import defaultdict

VENDORS = {
    'CASTING': 'VEN0000001;VEN0000003',
    'MACHINING': 'VEN0000002;VEN0000006',
    'HEAT_TREAT': 'VEN0000004',
    'GEAR': 'VEN0000005',
    'GRINDING': 'VEN0000006',
    'PLATING': 'VEN0000007',
    'HYDRAULIC': 'VEN0000008',
    'ELECTRICAL': 'VEN0000009',
    'SEAL': 'VEN0000010',
    'BOUGHTOUT': 'VEN0000011',
    'FASTENER': 'VEN0000012',
}

VENDOR_NAMES = {
    'VEN0000001': 'Aji Castings & Alloys Pvt Ltd',
    'VEN0000002': 'Apex Precision VMC Machinists',
    'VEN0000003': 'Shapar Heavy Foundry & Forging',
    'VEN0000004': 'Gujarat Heat Treaters & Nitriding',
    'VEN0000005': 'Maruti Precision Gear Works',
    'VEN0000006': 'Sardar Cylindrical & Surface Grinding',
    'VEN0000007': 'Hard Chrome Plating Corporation',
    'VEN0000008': 'Yuken India Hydraulics Pvt Ltd',
    'VEN0000009': 'Siemens Automation & Drives Ltd',
    'VEN0000010': 'Polymer Seal & O-Ring Industries',
    'VEN0000011': 'Kabra Extrusion & Machine Tools',
    'VEN0000012': 'National Fasteners & High Tensile Bolts',
}

wb = openpyxl.load_workbook('1-ERP BOM 90-200 TON.xlsx', data_only=True)

raw_rows = []
bom_structure = defaultdict(list)

for sname in wb.sheetnames:
    ws = wb[sname]
    max_r = ws.max_row
    max_c = ws.max_column
    
    row1 = [cell.value for cell in list(ws.iter_rows(min_row=1, max_row=1))[0]]
    row2 = [cell.value for cell in list(ws.iter_rows(min_row=2, max_row=2))[0]]
    
    if sname == 'Valve 90-200 TON':
        for col_start, blk, ton_model in [(1, 'HYDRAULIC VALVE GROUP - 90/110 TON', '90 / 110 TON'), (6, 'HYDRAULIC VALVE GROUP - 140/200 TON', '140 / 200 TON')]:
            sr_idx = 1
            for r in range(5, max_r + 1):
                sr = ws.cell(r, col_start+1).value
                desc = ws.cell(r, col_start+2).value
                qty = ws.cell(r, col_start+3).value
                sub_unit = ws.cell(r, col_start+4).value
                if desc and str(desc).strip():
                    name_clean = str(desc).strip()
                    code_slug = f"5000{abs(hash(name_clean)) % 900000 + 100000}"
                    row_dict = {
                        'sheet': sname,
                        'block': blk,
                        'row': r,
                        'sr_no': sr_idx,
                        'old_code': '',
                        'new_code': code_slug,
                        'name': name_clean,
                        'qty': float(qty) if qty and str(qty).replace('.','').isdigit() else 1.0,
                        'class': 'HY',
                        'sub_unit': str(sub_unit).strip() if sub_unit else ''
                    }
                    raw_rows.append(row_dict)
                    bom_structure[blk].append(row_dict)
                    sr_idx += 1
        continue

    col_idx = 0
    while col_idx < max_c:
        h1 = row1[col_idx]
        h2 = row2[col_idx] if col_idx < len(row2) else None
        
        if h2 and 'SR' in str(h2).upper():
            block_title = h1
            if not block_title:
                for b_idx in range(col_idx, -1, -1):
                    if row1[b_idx]:
                        block_title = row1[b_idx]
                        break
            
            subheaders = []
            for c in range(col_idx, max_c):
                val = row2[c]
                if c > col_idx and (val is None or (row1[c] and 'SR' not in str(val).upper() and row1[c] != block_title)):
                    break
                subheaders.append(str(val).strip().replace('\n', ' ') if val is not None else '')
                if 'CLASS' in subheaders[-1].upper():
                    break
            
            for r in range(3, max_r + 1):
                row_cells = [ws.cell(row=r, column=col_idx + 1 + offset).value for offset in range(len(subheaders))]
                sr_no = row_cells[0]
                if sr_no is None or str(sr_no).strip() == '':
                    continue
                try:
                    sr_num = int(float(str(sr_no).strip()))
                except:
                    continue
                
                row_dict = {'sheet': sname, 'block': str(block_title).strip(), 'row': r, 'sr_no': sr_num}
                for sh, cell_val in zip(subheaders, row_cells):
                    sh_clean = sh.upper()
                    if 'OLD' in sh_clean:
                        row_dict['old_code'] = str(cell_val).strip() if cell_val is not None else ''
                    elif 'NEW' in sh_clean or 'PART CODE' in sh_clean:
                        row_dict['new_code'] = str(cell_val).strip() if cell_val is not None else ''
                    elif 'DECRIPTION' in sh_clean or 'DESCRIPTION' in sh_clean or 'ITEM' in sh_clean:
                        row_dict['name'] = str(cell_val).strip() if cell_val is not None else ''
                    elif 'QTY' in sh_clean:
                        row_dict['qty'] = float(cell_val) if cell_val is not None and str(cell_val).strip() != '' else 1.0
                    elif 'CLASS' in sh_clean:
                        row_dict['class'] = str(cell_val).strip() if cell_val is not None else ''
                
                raw_rows.append(row_dict)
                bom_structure[str(block_title).strip()].append(row_dict)
            
            col_idx += max(1, len(subheaders))
        else:
            col_idx += 1

wb.close()

print(f"Loaded {len(raw_rows)} total item rows from 1-ERP BOM 90-200 TON.xlsx across {len(bom_structure)} BOM blocks.")

# Check Anomalies
code_to_names = defaultdict(set)
code_to_occs = defaultdict(list)
for item in raw_rows:
    code = item.get('new_code')
    name = item.get('name')
    if code and name:
        code_to_names[code].add(name)
        code_to_occs[code].append(item)

conflicting_codes = {k: v for k, v in code_to_names.items() if len(v) > 1}
print('\n=== 1. CONFLICTING PART CODES (Same Code -> Multiple Descriptions) ===')
if conflicting_codes:
    for code, names in conflicting_codes.items():
        print(f'Part Code: {code}')
        for name in names:
            occs = [f"{o['sheet']} | {o['block']} (Row {o['row']})" for o in code_to_occs[code] if o['name'] == name]
            print(f'  - Description: "{name}" in {occs}')
else:
    print('None found! (100% clean - no code has multiple descriptions)')

name_to_codes = defaultdict(set)
name_to_occs = defaultdict(list)
for item in raw_rows:
    code = item.get('new_code')
    name = item.get('name')
    if code and name:
        name_to_codes[name].add(code)
        name_to_occs[name].append(item)

conflicting_names = {k: v for k, v in name_to_codes.items() if len(v) > 1}
print(f'\n=== 2. CONFLICTING DESCRIPTIONS (Same Description -> Multiple Codes) === ({len(conflicting_names)} found)')
for name, codes in sorted(conflicting_names.items(), key=lambda x: x[0]):
    print(f'Description: "{name}"')
    for code in sorted(codes):
        occs = [f"{o['sheet']} | {o['block']} (Row {o['row']})" for o in name_to_occs[name] if o['new_code'] == code]
        print(f'  - Part Code: {code} in {occs}')
